from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph


DOWNLOADS = Path(r"C:\Users\ASUS\Downloads")
INPUTS = {
    "user": DOWNLOADS / "01_Manual_de_Usuario_Villegas2026_FINAL.docx",
    "technical": DOWNLOADS / "02_Documentacion_Tecnica_Desarrollador_Villegas2026_FINAL.docx",
    "install": DOWNLOADS / "03_Manual_Instalacion_Despliegue_Villegas2026_FINAL.docx",
}
OUTPUTS = {
    "user": DOWNLOADS / "01_Manual_de_Usuario_Villegas2026_REVISADO.docx",
    "technical": DOWNLOADS / "02_Documentacion_Tecnica_Desarrollador_Villegas2026_REVISADO.docx",
    "install": DOWNLOADS / "03_Manual_Instalacion_Despliegue_Villegas2026_REVISADO.docx",
}


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    paragraph.add_run(text)


def replace_in_paragraph(paragraph: Paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    set_paragraph_text(paragraph, paragraph.text.replace(old, new))
    return True


def all_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def replace_everywhere(doc: Document, replacements: dict[str, str]) -> None:
    for paragraph in all_paragraphs(doc):
        text = paragraph.text
        updated = text
        for old, new in replacements.items():
            updated = updated.replace(old, new)
        if updated != text:
            set_paragraph_text(paragraph, updated)


def table_text(table: Table) -> str:
    return "\n".join(cell.text for row in table.rows for cell in row.cells)


def remove_table(table: Table) -> None:
    table._element.getparent().remove(table._element)


def remove_tables_matching(doc: Document, pattern: str) -> None:
    rx = re.compile(pattern, re.IGNORECASE)
    for table in list(doc.tables):
        if rx.search(table_text(table)):
            remove_table(table)


def remove_paragraphs_matching(doc: Document, pattern: str) -> None:
    rx = re.compile(pattern, re.IGNORECASE)
    for paragraph in list(doc.paragraphs):
        if rx.search(paragraph.text):
            paragraph._element.getparent().remove(paragraph._element)


def remove_rows_matching(doc: Document, pattern: str) -> None:
    rx = re.compile(pattern, re.IGNORECASE)
    for table in doc.tables:
        for row in list(table.rows):
            text = " | ".join(cell.text for cell in row.cells)
            if rx.search(text):
                row._element.getparent().remove(row._element)


def insert_paragraph_before_table(table: Table, text: str, style: str | None = None) -> Paragraph:
    p = OxmlElement("w:p")
    table._tbl.addprevious(p)
    paragraph = Paragraph(p, table._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def insert_paragraph_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    p = OxmlElement("w:p")
    paragraph._p.addnext(p)
    new_paragraph = Paragraph(p, paragraph._parent)
    if style:
        new_paragraph.style = style
    new_paragraph.add_run(text)
    return new_paragraph


def update_control_and_history(doc: Document, description: str) -> None:
    for table in doc.tables:
        if not table.rows or len(table.rows[0].cells) < 2:
            continue
        if table.rows[0].cells[0].text.strip() != "Campo" or table.rows[0].cells[1].text.strip() != "Valor":
            continue
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            key = row.cells[0].text.strip()
            if key == "Fecha de revisión":
                set_paragraph_text(row.cells[1].paragraphs[0], "30/08/2026")
            elif key == "Versión":
                set_paragraph_text(row.cells[1].paragraphs[0], "1.2")

    replace_everywhere(doc, {
        "Fecha de revisión: 29/08/2026": "Fecha de revisión: 30/08/2026",
        "Versión del documento: 1.1": "Versión del documento: 1.2",
        "versión 1.1, revisada el 29/08/2026": "versión 1.2, revisada el 30/08/2026",
        "Versión documental 1.1 - 29/08/2026": "Versión documental 1.2 - 30/08/2026",
    })

    for table in doc.tables:
        if table.rows and len(table.rows[0].cells) >= 4:
            header = table.rows[0].cells
            if header[0].text.strip() == "Versión" and header[2].text.strip() == "Responsable" and header[3].text.strip() == "Descripción":
                set_paragraph_text(header[1].paragraphs[0], "Fecha")
                if not any(row.cells[0].text.strip() == "1.2" for row in table.rows[1:]):
                    row = table.add_row()
                    values = ["1.2", "30/08/2026", "Consorcios Villegas E.I.R.L.", description]
                    for cell, value in zip(row.cells, values):
                        set_paragraph_text(cell.paragraphs[0], value)
                break


def revise_user(doc: Document) -> None:
    replacements = {
        "Los catálogos soportan operaciones comerciales y permiten estandarizar valores usados en productos, documentos, pagos y comprobantes. Según las rutas revisadas se incluyen: unidades, tipos de afectación, tipos de documento, tipos de comprobante, tipos de cobranza, formas de pago, medios de pago, líneas, tipos de operación, certificados SUNAT, series de comprobantes y configuraciones.":
            "El sistema separa los catálogos operativos de la configuración administrativa. Las opciones visibles dependen de los permisos asignados al usuario; por ello no todos los perfiles verán los mismos submenús.",
        "1. Ingrese al catálogo que necesita administrar.":
            "1. En Configuración > Empresa y SUNAT se administra el certificado SUNAT cuando el usuario cuenta con el permiso correspondiente.",
        "2. Revise registros existentes antes de crear uno nuevo para evitar duplicados.":
            "2. En Configuración > Pagos y cobranza se encuentran Formas de Pago, Medios de Pago y Tipos de Cobranza.",
        "3. Use la acción de creación y complete los campos obligatorios.":
            "3. En Configuración > Comprobantes y documentos se encuentran Tipos de Operación, Tipos de Comprobante, Correlativos, Tipos de Documento y Tipos de Afectación.",
        "4. Edite únicamente cuando la modificación no afecte la interpretación histórica de operaciones existentes.":
            "4. En Configuración > Gastos y costos se administran categorías y tipos; en Parámetros se mantiene el costo del servicio de preparada.",
        "5. Desactive o elimine registros solo cuando el diseño del módulo lo permita y la operación sea segura.":
            "5. En Catálogos se encuentran Productos, Unidades y Líneas. La operación completa de Productos se explica en el capítulo 7. Antes de crear, editar o eliminar un registro, revise duplicados y confirme que el cambio no altere la interpretación histórica de operaciones existentes.",
        "13. Caja": "13. Caja: ingresos y reportes",
        "13.1. Pagos de caja": "13.1. Ingresos de caja",
        "1. Abra Caja - Pagos.": "1. Abra Caja y Provisionales > Ingresos Caja.",
        "2. Registre el egreso/pago con la información requerida.": "2. Seleccione Nuevo Ingreso.",
        "3. Revise monto, fecha, concepto y cuenta/medio antes de guardar.": "3. Complete Fecha, Monto, Caja destino (Principal, Depósito o Consorcio) y, si corresponde, un Comentario.",
        "4. Verifique la operación en el listado y en el reporte de caja.": "4. Guarde el registro. El usuario que realiza la operación se asigna automáticamente.",
        "13.2. Ingresos de caja": "13.2. Reporte de caja",
        "1. Abra Caja - Ingresos.": "1. Abra Caja y Provisionales > Reporte Caja.",
        "2. Registre el ingreso con los datos solicitados.": "2. Elija Reporte General de Caja o Reporte Detallado de Caja.",
        "3. Revise el monto y origen.": "3. Indique Fecha inicio y Fecha fin; la fecha final debe ser igual o posterior a la inicial.",
        "4. Guarde y verifique el registro.": "4. Seleccione Filtrar. En el reporte general también puede abrir la salida PDF.",
        "Cada recuadro identificado como “CAPTURA PENDIENTE MU-xx” indica la vista que debe incorporarse. Use datos de prueba o anonimice información personal. No incluya contraseñas, tokens, APP_KEY, datos bancarios ni secretos de configuración.":
            "Las capturas incluidas corresponden a vistas reales del sistema. Deben conservarse anonimizadas y no deben mostrar contraseñas, tokens, APP_KEY, datos bancarios ni secretos de configuración.",
        "Los recuadros ya están ubicados dentro del capítulo correspondiente. Sustituya cada recuadro por la captura o diagrama indicado y conserve el código como pie de figura para mantener trazabilidad.":
            "Las capturas incorporadas se encuentran dentro del capítulo correspondiente y conservan su código como pie de figura para mantener trazabilidad.",
        "Nota de control: las capturas identificadas como pendientes pueden incorporarse posteriormente sin alterar el contenido funcional aprobado, siempre que correspondan a la misma versión del sistema y no expongan datos personales ni credenciales.":
            "Nota de control: las capturas incorporadas forman parte de esta edición. Cualquier sustitución posterior debe corresponder a la misma versión funcional del sistema y no debe exponer datos personales ni credenciales.",
    }
    replace_everywhere(doc, replacements)

    remove_tables_matching(doc, r"CAPTURA PENDIENTE MU-(08|09|10|27|28)")
    remove_paragraphs_matching(doc, r"Figura MU-(08|09|10|27|28)\.")
    remove_rows_matching(doc, r"\bMU-(08|09|10|27|28)\b")

    # Add the actual edit/delete behavior to the active Caja Ingresos flow.
    target = next((p for p in doc.paragraphs if p.text.strip() == "4. Guarde el registro. El usuario que realiza la operación se asigna automáticamente."), None)
    if target is not None:
        insert_paragraph_after(target, "5. Según sus permisos, puede editar o eliminar un ingreso desde la columna Opciones.", target.style.name)


def revise_technical(doc: Document) -> None:
    replacements = {
        "Ingresos, pagos y reportes de caja.": "Ingresos y reportes general/detallado de caja.",
        "RF-020 | Caja | Registrar ingresos y pagos de caja.": "RF-020 | Caja | Registrar ingresos de caja y consultar reportes general y detallado. El flujo Caja - Pagos existe en código, pero no está habilitado en el menú activo.",
        "Registrar ingresos y pagos de caja.": "Registrar ingresos de caja y consultar reportes general y detallado; Caja - Pagos no está habilitado en el menú activo.",
        "/caja-pagos, /caja-ingresos": "/caja-ingresos, /reportes/caja",
        "Movimientos de caja.": "Ingresos de caja y reportes general/detallado. El código de /caja-pagos permanece como flujo no expuesto en el menú activo.",
        "RECOMENDADO - no verificado como implementación actual: No se encontró un manifiesto de despliegue, Dockerfile/docker-compose ni pipeline CI/CD. Esta sección propone una línea base basada en Laravel 12 y debe adaptarse al servidor real. [W3]":
            "VERIFICADO CON EL RESPONSABLE DEL SISTEMA: producción utiliza cPanel. El monolito Laravel se mantiene como una sola aplicación lógica, pero se separa físicamente: el contenido de public se publica en public_html y el resto del proyecto permanece en una carpeta privada hermana. No se encontraron Dockerfile ni pipeline CI/CD. [W3]",
        "16.1. Arquitectura de producción recomendada": "16.1. Arquitectura de producción en cPanel",
        "Laravel recomienda que el servidor web dirija las solicitudes al directorio public y no exponga la raíz del proyecto, donde podrían existir archivos de configuración sensibles. [W3]":
            "En cPanel, public_html contiene únicamente el contenido publicable de Laravel. El archivo public_html/index.php carga maintenance.php, vendor/autoload.php y bootstrap/app.php desde la carpeta privada del backend y ejecuta $app->usePublicPath(__DIR__) para que public_path() apunte a public_html. Esta adaptación es necesaria porque el proyecto usa public_path() para archivos y reportes. [W3]",
        "Los recuadros ya están ubicados dentro del capítulo correspondiente. Sustituya cada recuadro por la captura o diagrama indicado y conserve el código como pie de figura para mantener trazabilidad.":
            "Los diagramas DT-01, DT-02, DT-03, DT-04 y DT-06 a DT-11 pueden generarse con 02_Diagramas_Tecnicos_Villegas2026.puml. DT-05 continúa siendo una captura real del entorno de desarrollo, no un diagrama.",
        "Documentar servidor, PHP-FPM/Nginx, variables, permisos de storage, backups y rollback.":
            "Documentar cPanel, versión de PHP/Apache, variables, permisos de storage, backups y rollback.",
    }
    replace_everywhere(doc, replacements)

    # Word continued this ordered list from an earlier list in the document.
    # Render the release sequence with explicit 1-10 labels so it is stable.
    in_release_sequence = False
    release_step = 0
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "16.2. Secuencia de release recomendada":
            in_release_sequence = True
            continue
        if not in_release_sequence:
            continue
        if paragraph.style.name == "Code Block":
            break
        if paragraph.style.name == "List Number":
            release_step += 1
            original_text = paragraph.text.strip()
            if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
                paragraph._p.pPr.remove(paragraph._p.pPr.numPr)
            paragraph.style = doc.styles["Normal"]
            set_paragraph_text(paragraph, f"{release_step}. {original_text}")

    diagram_ids = ["DT-01", "DT-02", "DT-03", "DT-04", "DT-06", "DT-07", "DT-08", "DT-09", "DT-10", "DT-11"]
    for table in doc.tables:
        text = table_text(table)
        for diagram_id in diagram_ids:
            if diagram_id in text and "CAPTURA PENDIENTE" in text:
                for paragraph in (p for row in table.rows for cell in row.cells for p in cell.paragraphs):
                    replace_in_paragraph(paragraph, "CAPTURA PENDIENTE", "DIAGRAMA PENDIENTE")
                    replace_in_paragraph(
                        paragraph,
                        "Reemplazar este recuadro con una captura anonimizada del ambiente autorizado.",
                        "Generar desde 02_Diagramas_Tecnicos_Villegas2026.puml.",
                    )
                break

    for paragraph in doc.paragraphs:
        if any(f"Figura {diagram_id}." in paragraph.text for diagram_id in diagram_ids):
            text = re.sub(r"Fuente: captura del sistema Consorcios Villegas E\.I\.R\.L\.", "Fuente: código PlantUML verificado contra el proyecto.", paragraph.text)
            set_paragraph_text(paragraph, text)

    # Add the cPanel topology explanation immediately before DT-06.
    dt06 = next((t for t in doc.tables if "DT-06" in table_text(t) and "DIAGRAMA PENDIENTE" in table_text(t)), None)
    if dt06 is not None:
        insert_paragraph_before_table(
            dt06,
            "La carpeta privada contiene app, bootstrap, config, database, resources, routes, storage, vendor, artisan y .env. public_html contiene index.php, .htaccess, build y los demás assets copiados desde public; nunca debe contener .env, app, vendor ni la raíz completa del repositorio.",
            "Normal",
        )


def revise_install(doc: Document) -> None:
    replacements = {
        "En desarrollo, la aplicación se ejecuta como un monolito Laravel: el navegador consume rutas web, los controladores orquestan servicios de negocio, Eloquent persiste en MySQL y Blade/Vite entregan la interfaz. En producción, el repositorio no documenta una topología definitiva de servidor web; por ello la figura siguiente debe adaptarse al ambiente real.":
            "En desarrollo, la aplicación se ejecuta como un monolito Laravel. En producción continúa siendo un monolito lógico, pero cPanel exige una separación física: el contenido publicable se coloca en public_html y todo lo demás permanece en una carpeta privada hermana. El navegador solo accede a public_html; index.php inicia el backend privado, que ejecuta controladores y servicios, usa Eloquent/MySQL y renderiza Blade.",
        "Servidor web configurado para usar la carpeta public como document root.":
            "Cuenta cPanel con PHP 8.2+; public_html debe contener solo el contenido de public y la carpeta privada del backend debe quedar fuera del alcance web.",
        "13. Servidor web y PHP": "13. cPanel: PHP, backend privado y public_html",
        "Laravel debe publicarse desde la carpeta public. El siguiente bloque es un EJEMPLO de Nginx, no evidencia de que Nginx sea el servidor utilizado actualmente. Ajuste dominio, socket PHP-FPM, límites y cabeceras al ambiente real.":
            "En cPanel no se publica la raíz completa del repositorio. Copie el contenido de public en public_html y mantenga el resto en /home/<cuenta>/<carpeta_backend>. Después adapte public_html/index.php como se muestra; sustituya <carpeta_backend> por el nombre real de la carpeta privada.",
        "No copie este bloque sin validarlo. Si el ambiente real usa Apache, contenedores, IIS u otro servicio administrado, documente esa topología en OPS-01 y sustituya este ejemplo.":
            "Conserve public_html/.htaccess desde public/.htaccess. No coloque .env, app, bootstrap, config, database, resources, routes, storage ni vendor dentro de public_html. El nombre real de <carpeta_backend> debe validarse en la cuenta cPanel.",
        "Laravel necesita escritura en storage y bootstrap/cache. La identidad exacta del usuario/grupo depende del servidor. Evite chmod 777. Ejemplo conceptual:":
            "Laravel necesita escritura en <carpeta_backend>/storage y <carpeta_backend>/bootstrap/cache. Use el Administrador de archivos o la Terminal de cPanel y evite chmod 777. Ejemplo conceptual desde /home/<cuenta>:",
        "sudo chown -R <usuario_php>:<grupo_php> storage bootstrap/cache\nsudo chmod -R u+rwX,g+rwX storage bootstrap/cache":
            "chmod -R u+rwX,g+rwX <carpeta_backend>/storage <carpeta_backend>/bootstrap/cache",
        "Si el sistema guarda archivos públicos mediante el filesystem de Laravel, verifique también el enlace public/storage según la configuración real. No se debe asumir su necesidad si el módulo no lo utiliza.":
            "El index.php adaptado ejecuta $app->usePublicPath(__DIR__), de modo que public_path() resuelve a public_html. Esto es importante porque el proyecto guarda imágenes de productos mediante public_path('uploads/productos') y algunos reportes consultan archivos públicos. Si se habilita el disco public de Laravel, cree y valide aparte el enlace public_html/storage hacia <carpeta_backend>/storage/app/public.",
        "5. Actualizar el código al commit aprobado.": "5. Actualizar el código de la carpeta privada al commit aprobado, sin copiar .env ni datos operativos.",
        "7. Instalar/compilar assets frontend.": "7. Compilar los assets en un entorno con Node.js y sincronizar el contenido actualizado de public con public_html; como mínimo, verificar public_html/build/manifest.json.",
        "10. Levantar la aplicación y ejecutar smoke tests.": "10. Confirmar que public_html/index.php conserva las rutas a <carpeta_backend> y ejecutar smoke tests por HTTPS.",
        "# Ejemplo de secuencia; adaptar a su CI/CD": "# Ejecutar desde /home/<cuenta>/<carpeta_backend>",
        "git checkout <commit-aprobado>\ncomposer install --no-dev --optimize-autoloader\nnpm ci\nnpm run build\nphp artisan migrate --force\nphp artisan optimize:clear":
            "git checkout <commit-aprobado>\ncomposer install --no-dev --optimize-autoloader\n# Compilar localmente si cPanel no ofrece Node.js:\n# npm ci && npm run build\n# Sincronizar el contenido de public/ con /home/<cuenta>/public_html/\nphp artisan migrate --force\nphp artisan optimize:clear\nphp artisan optimize",
        "php artisan down/up se presenta como mecanismo recomendado para una ventana controlada. Si la infraestructura usa balanceadores, blue/green o despliegues sin interrupción, reemplace esta secuencia por el procedimiento real.":
            "php artisan down/up se ejecuta desde la carpeta privada. Antes de php artisan up, confirme que public_html contiene el .htaccess y los assets del release, y que index.php apunta a la carpeta privada correcta.",
        "Los recuadros ya están ubicados dentro del capítulo correspondiente. Sustituya cada recuadro por la captura o diagrama indicado y conserve el código como pie de figura para mantener trazabilidad.":
            "OPS-01 puede generarse con 03_Diagrama_Despliegue_cPanel_Villegas2026.puml. OPS-02 a OPS-07 son evidencias reales del proceso y no deben reemplazarse por diagramas.",
        "No se declara en el repositorio una versión exacta de MySQL, Node.js, Nginx o Apache. Defina y registre las versiones que se usen realmente en producción antes de aprobar este manual.":
            "El despliegue utiliza cPanel/Apache, pero el repositorio no declara las versiones exactas de MySQL, Node.js, Apache ni la configuración PHP del hosting. Registre las versiones reales de la cuenta cPanel antes de aprobar el manual.",
    }
    replace_everywhere(doc, replacements)

    # Replace the former Nginx example with the real cPanel front controller adaptation.
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name == "Code Block" and paragraph.text.lstrip().startswith("server {"):
            set_paragraph_text(
                paragraph,
                "<?php\n\nuse Illuminate\\Foundation\\Application;\nuse Illuminate\\Http\\Request;\n\ndefine('LARAVEL_START', microtime(true));\n\nif (file_exists($maintenance = __DIR__.'/../<carpeta_backend>/storage/framework/maintenance.php')) {\n    require $maintenance;\n}\n\nrequire __DIR__.'/../<carpeta_backend>/vendor/autoload.php';\n\n/** @var Application $app */\n$app = require_once __DIR__.'/../<carpeta_backend>/bootstrap/app.php';\n$app->usePublicPath(__DIR__);\n\n$app->handleRequest(Request::capture());",
            )
            break

    ops01 = next((t for t in doc.tables if "OPS-01" in table_text(t)), None)
    if ops01 is not None:
        for paragraph in (p for row in ops01.rows for cell in row.cells for p in cell.paragraphs):
            replace_in_paragraph(paragraph, "CAPTURA PENDIENTE", "DIAGRAMA PENDIENTE")
            replace_in_paragraph(paragraph, "Diagrama de despliegue real", "Diagrama de despliegue cPanel")
            replace_in_paragraph(
                paragraph,
                "Reemplazar este recuadro con una captura anonimizada del ambiente autorizado.",
                "Generar desde 03_Diagrama_Despliegue_cPanel_Villegas2026.puml.",
            )
        insert_paragraph_before_table(
            ops01,
            "/home/<cuenta>/\n├── <carpeta_backend>/   # app, bootstrap, config, database, resources, routes, storage, vendor, artisan y .env\n└── public_html/         # contenido de public: index.php adaptado, .htaccess, build, assets, css, js y archivos públicos",
            "Code Block",
        )

    # Update the evidence checklist wording for OPS-01.
    for table in doc.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() == "OPS-01":
                if len(row.cells) > 1:
                    set_paragraph_text(row.cells[1].paragraphs[0], "Diagrama de despliegue cPanel")
                if len(row.cells) > 2:
                    set_paragraph_text(row.cells[2].paragraphs[0], "public_html separado de la carpeta privada del backend y conexión con MySQL.")


def main() -> None:
    jobs = [
        ("user", revise_user, "Corrección de Catálogos y del flujo activo de Caja; conservación de capturas incorporadas."),
        ("technical", revise_technical, "Ajuste de Caja, diagramas PlantUML y arquitectura productiva cPanel."),
        ("install", revise_install, "Adecuación del despliegue real cPanel con backend privado y public_html."),
    ]
    for key, editor, history in jobs:
        doc = Document(str(INPUTS[key]))
        editor(doc)
        update_control_and_history(doc, history)
        doc.save(str(OUTPUTS[key]))
        print(OUTPUTS[key])


if __name__ == "__main__":
    main()
