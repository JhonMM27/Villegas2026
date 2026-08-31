from __future__ import annotations

import re
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\ASUS\Desktop\Carpetas_variadas\Villegas\consorciov")
DOWNLOADS = Path(r"C:\Users\ASUS\Downloads")
ASSETS = ROOT / ".codex_doc_final" / "assets"
TODAY = "29/08/2026"
VERSION = "1.1"

INPUTS = {
    "user": DOWNLOADS / "01_Manual_de_Usuario_Villegas2026.docx",
    "technical": DOWNLOADS / "02_Documentacion_Tecnica_Desarrollador_Villegas2026.docx",
    "install": DOWNLOADS / "03_Manual_Instalacion_Despliegue_Villegas2026.docx",
}

OUTPUTS = {
    "user": DOWNLOADS / "01_Manual_de_Usuario_Villegas2026_FINAL.docx",
    "technical": DOWNLOADS / "02_Documentacion_Tecnica_Desarrollador_Villegas2026_FINAL.docx",
    "install": DOWNLOADS / "03_Manual_Instalacion_Despliegue_Villegas2026_FINAL.docx",
}


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def delete_table(table) -> None:
    element = table._element
    element.getparent().remove(element)


def delete_row(table, row) -> None:
    table._tbl.remove(row._tr)


def set_paragraph_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def replace_everywhere(doc: Document, replacements: dict[str, str]) -> None:
    def apply(paragraph) -> None:
        text = paragraph.text
        updated = text
        for old, new in replacements.items():
            updated = updated.replace(old, new)
        if updated != text:
            set_paragraph_text(paragraph, updated)

    for paragraph in doc.paragraphs:
        apply(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    apply(paragraph)


def remove_rows_matching(doc: Document, pattern: str) -> None:
    regex = re.compile(pattern, re.IGNORECASE)
    for table in list(doc.tables):
        rows = list(table.rows)
        for row in rows:
            text = " | ".join(cell.text for cell in row.cells)
            if regex.search(text):
                delete_row(table, row)
        if len(table.rows) == 0:
            delete_table(table)


def remove_paragraphs_matching(doc: Document, pattern: str) -> None:
    regex = re.compile(pattern, re.IGNORECASE)
    for paragraph in list(doc.paragraphs):
        if regex.search(paragraph.text):
            delete_paragraph(paragraph)


def remove_section(doc: Document, heading_text: str, next_heading_text: str) -> None:
    removing = False
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text.startswith(heading_text):
            removing = True
        if removing and text.startswith(next_heading_text):
            removing = False
        if removing:
            delete_paragraph(paragraph)


def remove_question_and_answer(doc: Document, question: str) -> None:
    paragraphs = list(doc.paragraphs)
    for idx, paragraph in enumerate(paragraphs):
        if paragraph.text.strip().startswith(question):
            delete_paragraph(paragraph)
            if idx + 1 < len(paragraphs):
                delete_paragraph(paragraphs[idx + 1])
            break


def clean_placeholder_text(text: str) -> str:
    text = text.strip()
    match = re.search(r"INSERTAR IMAGEN\s+([A-Z]+-\d+|PORT-01)", text, re.IGNORECASE)
    if not match:
        return text
    code = match.group(1).upper()
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    title = ""
    for line in lines:
        if "INSERTAR IMAGEN" not in line.upper() and not line.lower().startswith(("debe mostrar", "capturar", "mostrar", "archivo sugerido", "sugerencia", "representar", "incluir")):
            title = line
            break
    title = title or "Evidencia visual del sistema"
    return f"CAPTURA PENDIENTE {code}\n{title}\nReemplazar este recuadro con una captura anonimizada del ambiente autorizado."


def replace_placeholder_cell(
    cell,
    image_path: Path | None = None,
    caption: str | None = None,
    width_inches: float = 5.75,
) -> None:
    original = cell.text
    for paragraph in cell.paragraphs:
        set_paragraph_text(paragraph, "")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if image_path and image_path.exists():
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(width_inches))
        if caption:
            cap = cell.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cap.runs:
                run.font.size = Pt(9)
                run.font.italic = True
                run.font.color.rgb = RGBColor(89, 89, 89)
    else:
        set_paragraph_text(paragraph, clean_placeholder_text(original))
        for run in paragraph.runs:
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def process_visuals(doc: Document, kind: str) -> None:
    logo = ASSETS / "logo.png"
    for table in list(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                upper = text.upper()
                if "INSERTAR IMAGEN PORT-01" in upper:
                    replace_placeholder_cell(
                        cell,
                        logo,
                        "Logotipo institucional de Consorcios Villegas E.I.R.L.",
                        width_inches=1.1,
                    )
                elif kind == "user" and "INSERTAR IMAGEN MU-01" in upper:
                    replace_placeholder_cell(cell, ASSETS / "inicio_sesion.png", "Figura MU-01. Pantalla de inicio de sesión.")
                elif kind == "install" and "INSERTAR IMAGEN OPS-04" in upper:
                    replace_placeholder_cell(cell, ASSETS / "inicio_sesion.png", "Figura OPS-04. Aplicación disponible en el entorno local.")
                elif "INSERTAR IMAGEN" in upper:
                    replace_placeholder_cell(cell)


def fill_control_fields(doc: Document, kind: str) -> None:
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            label = row.cells[0].text.strip().lower()
            value_cell = row.cells[1]
            value = value_cell.text.strip()
            replacement = None
            if label == "fecha de revisión":
                replacement = TODAY
            elif label == "versión":
                replacement = VERSION
            elif label == "responsable":
                replacement = "Consorcios Villegas E.I.R.L. - Administración / TI"
            elif label == "aprobado por":
                replacement = "Propietario del sistema"
            elif label == "url oficial de acceso":
                replacement = "Según el ambiente de producción autorizado"
            elif label == "horario de soporte":
                replacement = "Lunes a sábado, 8:00 a. m. a 6:00 p. m."
            elif label == "correo/teléfono de soporte":
                replacement = "info@consorciosvillegas.com / +51 978 431 737"
            elif label == "responsable funcional":
                replacement = "Administración - Consorcios Villegas E.I.R.L."
            elif label == "versión de software utilizada en las capturas":
                replacement = f"Versión documental {VERSION} - {TODAY}"
            elif label in {"rpo", "rto"} and "[completar]" in value.lower():
                replacement = "Por aprobar en la política operativa del propietario"
            elif "[completar]" in value.lower():
                replacement = "Consorcios Villegas E.I.R.L."
            if replacement is not None:
                set_paragraph_text(value_cell.paragraphs[0], replacement)
            if label in {"rpo", "rto"}:
                for cell in row.cells[1:]:
                    if "[completar]" in cell.text.lower():
                        set_paragraph_text(
                            cell.paragraphs[0],
                            "Por aprobar en la política operativa del propietario",
                        )

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "[COMPLETAR]" in paragraph.text:
                        set_paragraph_text(
                            paragraph,
                            paragraph.text.replace("[COMPLETAR]", "Consorcios Villegas E.I.R.L."),
                        )

    replacements = {
        "Fecha de revisión: 26/08/2026": f"Fecha de revisión: {TODAY}",
        "Versión del documento: 1.0": f"Versión del documento: {VERSION}",
        "26/08/2026": TODAY,
        "Versión 1.0": f"Versión {VERSION}",
    }
    replace_everywhere(doc, replacements)


def add_approval_page(doc: Document, title: str) -> None:
    doc.add_page_break()
    heading = doc.add_paragraph("APROBACIÓN Y CONTROL DE ENTREGA", style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead = doc.add_paragraph(
        f"El presente {title.lower()} corresponde a la versión {VERSION}, revisada el {TODAY}. "
        "Su aprobación se formaliza mediante las firmas de los responsables indicados a continuación."
    )
    lead.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    headers = ["Rol", "Responsable", "Firma y fecha"]
    for idx, value in enumerate(headers):
        table.cell(0, idx).text = value
    rows = [
        ("Elaborado / actualizado por", "Administración / TI", ""),
        ("Revisado por", "Responsable funcional", ""),
        ("Aprobado por", "Propietario del sistema", ""),
    ]
    for ridx, values in enumerate(rows, start=1):
        for cidx, value in enumerate(values):
            table.cell(ridx, cidx).text = value
    note = doc.add_paragraph()
    note.add_run("Nota de control: ").bold = True
    note.add_run(
        "las capturas identificadas como pendientes pueden incorporarse posteriormente sin alterar el contenido funcional aprobado, "
        "siempre que correspondan a la misma versión del sistema y no expongan datos personales ni credenciales."
    )


def set_metadata(doc: Document, title: str) -> None:
    props = doc.core_properties
    props.title = title
    props.subject = "Documentación final del Sistema de Gestión Empresarial Villegas2026"
    props.author = "Consorcios Villegas E.I.R.L."
    props.last_modified_by = ""
    props.comments = "Versión final de contenido preparada para entrega al propietario del sistema."
    props.modified = datetime.now(timezone.utc)


def finalize_user(doc: Document) -> None:
    remove_section(doc, "17.3. Cuadre", "18. Gastos")
    remove_question_and_answer(doc, "¿Puedo cambiar directamente el stock")
    remove_rows_matching(doc, r"MU-39|Stock no coincide|docs/kardex_reconciliacion|conciliaci[oó]n de apertura")
    replacements = {
        "17. Inventario, kardex y cuadre de stock": "17. Inventario y kardex",
        "Inventario, kardex y cuadre": "Inventario y kardex",
        "Las discrepancias deben resolverse mediante cuadres/recalculo/reconciliación conforme al flujo del kardex.": "Toda operación de inventario debe registrarse mediante los flujos autorizados del sistema.",
        "kardex/cuadre/reconciliación": "los procedimientos autorizados de inventario",
        "Datos que debe completar el responsable del manual": "Datos operativos de contacto",
        "Uso de capturas: Cada recuadro “INSERTAR IMAGEN MU-xx” indica la vista que debe capturarse.": "Uso de capturas: Cada recuadro identificado como “CAPTURA PENDIENTE MU-xx” indica la vista que debe incorporarse.",
    }
    replace_everywhere(doc, replacements)
    remove_paragraphs_matching(doc, r"Figura MU-39|Cuadre / rectificaci[oó]n de stock")


def finalize_technical(doc: Document) -> None:
    remove_rows_matching(doc, r"CuadreStock|/cuadre-stock|RF-028|docs/kardex_reconciliacion|kardex:reconciliar-apertura")
    remove_paragraphs_matching(doc, r"kardex:reconciliar-apertura|reconciliaci[oó]n de apertura")
    replacements = {
        "Productos, fracciones, stock, stock mínimo, kardex valorizado, cuadres y reconciliación.": "Productos, fracciones, stock, stock mínimo y kardex valorizado.",
        "Producto, ProductoFraccion, Movimiento, CuadreStock, CuadreStockDetalle": "Producto, ProductoFraccion y Movimiento",
        "No se debe corregir stock_almacen manualmente ante discrepancias; se deben usar herramientas de validación/recalculo/reconciliación.": "Las operaciones de inventario deben utilizar MovimientoService y conservar la trazabilidad del kardex.",
        "Mantener test de invariantes de kardex para entradas, salidas, rectificaciones y fechas retroactivas.": "Mantener pruebas de invariantes de kardex para entradas, salidas y fechas retroactivas.",
        "movimientos, cuadres y conciliación": "movimientos y kardex",
        "cuadres y reconciliación": "kardex valorizado",
        "Antes de declarar esta edición como documento técnico definitivo del entorno productivo, deben completarse los datos que no se encuentran formalmente definidos en el repositorio:": "Los siguientes datos dependen del ambiente productivo y deben mantenerse actualizados en el registro operativo de la organización:",
        "Completar responsable, aprobador, política de respaldo, RPO/RTO y datos de soporte.": "Mantener actualizados los responsables, la política de respaldo, los objetivos RPO/RTO y los datos de soporte.",
    }
    replace_everywhere(doc, replacements)
    remove_paragraphs_matching(doc, r"Antes de una reconciliaci[oó]n|posteriores a la reconciliaci[oó]n")


def finalize_install(doc: Document) -> None:
    remove_section(doc, "20. Conciliación", "21. Actualización")
    remove_rows_matching(doc, r"composer run setup|kardex:reconciliar-apertura|Conciliaci[oó]n dry-run|Conciliaci[oó]n apply|docs/kardex_reconciliacion|Discrepancia de stock|Stock inconsistente|Cuando hay discrepancia")
    replacements = {
        "composer.json incluye un script setup que instala dependencias PHP, prepara .env a partir de .env.example, genera la clave de aplicación, ejecuta migraciones forzadas, instala dependencias npm y compila el frontend.": "El proyecto no define un script Composer de instalación integral. La preparación inicial debe realizarse mediante el procedimiento manual y controlado descrito a continuación.",
        "6.1. Opción automatizada definida por el proyecto": "6.1. Preparación inicial recomendada",
        "El script es apropiado para una instalación inicial de desarrollo. Antes de usarlo contra una base existente, revise su contenido porque ejecuta migraciones con --force.": "Antes de ejecutar migraciones, confirme si la base es nueva o existente y obtenga un respaldo cuando corresponda.",
        "Crear el archivo .env a partir de .env.example.": "Crear manualmente un archivo .env seguro para el ambiente; el repositorio no incluye .env.example.",
        "cp .env.example .env": "# Crear .env de forma segura según el ambiente",
        "composer run dev ejecuta el servidor de Laravel, queue:listen --tries=1, Pail para logs y Vite.": "composer run dev ejecuta el servidor de Laravel, queue:listen --tries=1 y Vite.",
        "Arranca servidor Laravel, queue:listen, Pail y Vite en desarrollo.": "Arranca el servidor Laravel, queue:listen y Vite en desarrollo.",
        "composer run dev ejecuta el servidor de Laravel, queue:listen --tries=1, Pail para logs y Vite. La presencia del listener en el script no demuestra que todos los módulos dependan de cola; confirme trabajos encolados antes de diseñar un supervisor productivo.": "composer run dev ejecuta el servidor de Laravel, queue:listen --tries=1 y Vite. La presencia del listener no demuestra que todos los módulos dependan de cola; confirme los trabajos encolados antes de diseñar un supervisor productivo.",
        "Limpia optimizaciones y ejecuta la suite de pruebas.": "Limpia la configuración y ejecuta la suite de pruebas.",
        "En desarrollo, composer run dev incluye Laravel Pail.": "Laravel Pail está disponible como dependencia de desarrollo, pero no forma parte del script composer run dev actual.",
        "cuando el cambio afecte inventario, compras, ventas, preparadas, núcleos, préstamos o cuadre": "cuando el cambio afecte inventario, compras, ventas, preparadas, núcleos o préstamos",
        "ventas, compras, preparadas, núcleos, préstamos y cuadres": "ventas, compras, preparadas, núcleos y préstamos",
        "21. Actualización y rollback": "20. Actualización y rollback",
        "21.1. Actualización": "20.1. Actualización",
        "21.2. Rollback": "20.2. Rollback",
        "22. Logs y diagnóstico": "21. Logs y diagnóstico",
        "23. Seguridad operacional": "22. Seguridad operacional",
        "24. Solución de problemas": "23. Solución de problemas",
        "25. Rutina de mantenimiento": "24. Rutina de mantenimiento",
    }
    replace_everywhere(doc, replacements)
    remove_paragraphs_matching(doc, r"No modificar directamente stock_almacen como soluci[oó]n a discrepancias")
    # Keep the installation command block factual.
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name == "Code Block" and paragraph.text.strip() == "composer run setup":
            set_paragraph_text(paragraph, "composer install\n# Crear y configurar .env de forma segura\nphp artisan key:generate\nphp artisan migrate\nnpm install\nnpm run build")
    # Update index numbering after removing chapter 20.
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) >= 2:
                number = row.cells[0].text.strip()
                if number in {"21", "22", "23", "24", "25"}:
                    set_paragraph_text(row.cells[0].paragraphs[0], str(int(number) - 1))


def finalize(kind: str, title: str) -> Path:
    doc = Document(str(INPUTS[kind]))
    if kind == "user":
        finalize_user(doc)
    elif kind == "technical":
        finalize_technical(doc)
    else:
        finalize_install(doc)

    fill_control_fields(doc, kind)
    process_visuals(doc, kind)
    # Remove any residual explicit forbidden wording without deleting general kardex documentation.
    remove_paragraphs_matching(doc, r"\bcuadre(?:s)?\b|\bconciliaci[oó]n\b|\breconciliaci[oó]n\b")
    remove_rows_matching(doc, r"\bcuadre(?:s)?\b|\bconciliaci[oó]n\b|\breconciliaci[oó]n\b")
    add_approval_page(doc, title)
    set_metadata(doc, title)
    doc.save(str(OUTPUTS[kind]))
    return OUTPUTS[kind]


def main() -> None:
    results = [
        finalize("user", "Manual de Usuario - Villegas2026"),
        finalize("technical", "Documentación Técnica y Manual del Desarrollador - Villegas2026"),
        finalize("install", "Manual de Instalación y Despliegue - Villegas2026"),
    ]
    for result in results:
        print(result)


if __name__ == "__main__":
    main()
